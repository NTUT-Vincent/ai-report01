from __future__ import annotations

import hashlib
import re
import uuid
from pathlib import Path
from typing import Any

import yaml
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from pydantic_settings import BaseSettings
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, sessionmaker


class Settings(BaseSettings):
    database_url: str = "postgresql://okf:okf@localhost:5432/okf"
    raw_data_dir: str = "./data/raw"
    max_upload_size_mb: int = 20
    cors_origins: str = "http://localhost:5173"

    class Config:
        env_file = ".env"


settings = Settings()
engine = create_engine(settings.database_url, future=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
RAW_DIR = Path(settings.raw_data_dir)
RAW_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="OKF Knowledge Factory MVP", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in settings.cors_origins.split(",") if o.strip()],
    allow_methods=["*"],
    allow_headers=["*"],
)


def db() -> Session:
    s = SessionLocal()
    try:
        yield s
    finally:
        s.close()


class ValidateOKFRequest(BaseModel):
    okf_json: dict[str, Any]


class ApproveOKFRequest(BaseModel):
    okf_json: dict[str, Any]
    approved_by: str = Field(min_length=1)


class RejectOKFRequest(BaseModel):
    reason: str = Field(min_length=1)


class EntityReviewAction(BaseModel):
    result_id: str
    action: str
    entity_id: str | None = None
    canonical_name: str | None = None
    entity_type: str | None = None
    alias: str | None = None


class EntityReviewRequest(BaseModel):
    reviewed_by: str = Field(min_length=1)
    actions: list[EntityReviewAction]


def load_source() -> dict[str, Any]:
    with open("source.yaml", "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_name(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", name).strip("._")
    return clean or "document"


def set_status(s: Session, doc_id: str, status: str, error: str | None = None) -> None:
    s.execute(text("UPDATE documents SET status=:st, error_message=:err, updated_at=NOW() WHERE document_id=:id"), {"st": status, "err": error, "id": doc_id})
    s.commit()


def save_artifact(s: Session, doc_id: str, typ: str, content: dict[str, Any]) -> None:
    s.execute(
        text("INSERT INTO artifacts (artifact_id, document_id, artifact_type, content_json, created_at) VALUES (:id,:doc,:typ,CAST(:content AS JSONB),NOW())"),
        {"id": str(uuid.uuid4()), "doc": doc_id, "typ": typ, "content": to_json(content)},
    )
    s.commit()


def latest_artifact(s: Session, doc_id: str, typ: str) -> dict[str, Any]:
    row = s.execute(
        text("SELECT content_json FROM artifacts WHERE document_id=:doc AND artifact_type=:typ ORDER BY created_at DESC LIMIT 1"),
        {"doc": doc_id, "typ": typ},
    ).mappings().first()
    if not row:
        raise HTTPException(400, f"missing artifact: {typ}")
    return row["content_json"]


def to_json(obj: dict[str, Any]) -> str:
    import json
    return json.dumps(obj, ensure_ascii=False)


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/db/init")
def init_db(s: Session = Depends(db)) -> dict[str, str]:
    s.execute(text("CREATE EXTENSION IF NOT EXISTS vector"))
    s.execute(text("""
    CREATE TABLE IF NOT EXISTS documents (
      document_id UUID PRIMARY KEY,
      source_id TEXT NOT NULL,
      file_name TEXT NOT NULL,
      file_type TEXT NOT NULL,
      file_path TEXT NOT NULL,
      file_hash TEXT NOT NULL,
      status TEXT NOT NULL,
      error_message TEXT,
      created_at TIMESTAMP DEFAULT NOW(),
      updated_at TIMESTAMP DEFAULT NOW()
    );
    CREATE INDEX IF NOT EXISTS ix_documents_hash ON documents(file_hash);
    CREATE TABLE IF NOT EXISTS artifacts (
      artifact_id UUID PRIMARY KEY,
      document_id UUID REFERENCES documents(document_id) ON DELETE CASCADE,
      artifact_type TEXT NOT NULL,
      content_json JSONB,
      file_path TEXT,
      created_at TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS entities (
      entity_id UUID PRIMARY KEY,
      canonical_name TEXT NOT NULL,
      entity_type TEXT NOT NULL,
      description TEXT,
      created_at TIMESTAMP DEFAULT NOW(),
      updated_at TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS entity_aliases (
      alias_id UUID PRIMARY KEY,
      entity_id UUID REFERENCES entities(entity_id) ON DELETE CASCADE,
      alias TEXT NOT NULL,
      source_document_id UUID REFERENCES documents(document_id) ON DELETE SET NULL,
      confidence FLOAT,
      approved_by TEXT,
      approved_at TIMESTAMP,
      created_at TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS entity_resolution_results (
      result_id UUID PRIMARY KEY,
      document_id UUID REFERENCES documents(document_id) ON DELETE CASCADE,
      mention TEXT NOT NULL,
      decision TEXT NOT NULL,
      suggested_entity_id UUID,
      suggested_canonical_name TEXT,
      entity_type TEXT,
      confidence FLOAT NOT NULL,
      reason TEXT,
      requires_review BOOLEAN DEFAULT TRUE,
      review_status TEXT DEFAULT 'pending',
      created_at TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS okf_documents (
      okf_id UUID PRIMARY KEY,
      document_id UUID REFERENCES documents(document_id) ON DELETE CASCADE,
      schema_version TEXT NOT NULL,
      title TEXT NOT NULL,
      confidentiality TEXT NOT NULL,
      summary TEXT,
      okf_json JSONB NOT NULL,
      review_status TEXT NOT NULL,
      approved_by TEXT,
      approved_at TIMESTAMP,
      created_at TIMESTAMP DEFAULT NOW(),
      updated_at TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS okf_chunks (
      chunk_id UUID PRIMARY KEY,
      okf_id UUID REFERENCES okf_documents(okf_id) ON DELETE CASCADE,
      chunk_type TEXT NOT NULL,
      content TEXT NOT NULL,
      metadata_json JSONB,
      embedding vector(1536),
      created_at TIMESTAMP DEFAULT NOW()
    );
    """))
    s.commit()
    seed_entities(s)
    return {"status": "initialized"}


def seed_entities(s: Session) -> None:
    seeds = [
        ("Intelligent Manufacturing Center", "organization", ["CIM", "IMC", "智慧製造中心"]),
        ("equipment", "equipment", ["機台", "tool", "equipment"]),
        ("SOP", "document_type", ["標準作業程序", "Standard Operating Procedure", "SOP"]),
        ("alarm code", "domain_term", ["alarm", "警報碼", "alarm code"]),
    ]
    for name, typ, aliases in seeds:
        row = s.execute(text("SELECT entity_id FROM entities WHERE lower(canonical_name)=lower(:name)"), {"name": name}).first()
        eid = str(row[0]) if row else str(uuid.uuid4())
        if not row:
            s.execute(text("INSERT INTO entities(entity_id, canonical_name, entity_type) VALUES (:id,:name,:typ)"), {"id": eid, "name": name, "typ": typ})
        for alias in aliases:
            exists = s.execute(text("SELECT 1 FROM entity_aliases WHERE lower(alias)=lower(:a) AND entity_id=:e"), {"a": alias, "e": eid}).first()
            if not exists:
                s.execute(text("INSERT INTO entity_aliases(alias_id, entity_id, alias, confidence, approved_by, approved_at) VALUES (:id,:e,:a,1,'seed',NOW())"), {"id": str(uuid.uuid4()), "e": eid, "a": alias})
    s.commit()


@app.get("/config/source")
def get_source() -> dict[str, Any]:
    return load_source()


@app.post("/documents/upload")
async def upload(file: UploadFile = File(...), s: Session = Depends(db)) -> dict[str, Any]:
    source = load_source()
    ext = Path(file.filename or "").suffix.lower().lstrip(".")
    if ext not in source["parser"]["supported_types"]:
        raise HTTPException(400, f"unsupported file type: {ext}")
    data = await file.read()
    if len(data) > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(400, "file too large")
    digest = sha256_bytes(data)
    matched = s.execute(text("SELECT document_id FROM documents WHERE file_hash=:h LIMIT 1"), {"h": digest}).first()
    doc_id = str(uuid.uuid4())
    filename = f"{doc_id}_{safe_name(file.filename or 'document.' + ext)}"
    path = RAW_DIR / filename
    path.write_bytes(data)
    status = "DUPLICATE" if matched else "RECEIVED"
    s.execute(text("INSERT INTO documents(document_id, source_id, file_name, file_type, file_path, file_hash, status) VALUES (:id,:src,:fn,:ft,:fp,:fh,:st)"), {"id": doc_id, "src": source["source_id"], "fn": file.filename, "ft": ext, "fp": str(path), "fh": digest, "st": status})
    s.commit()
    return {"document_id": doc_id, "status": status, "matched_document_id": str(matched[0]) if matched else None}


@app.get("/documents")
def documents(status: str | None = None, s: Session = Depends(db)) -> list[dict[str, Any]]:
    q = "SELECT * FROM documents" + (" WHERE status=:st" if status else "") + " ORDER BY created_at DESC"
    return [dict(r) for r in s.execute(text(q), {"st": status}).mappings().all()]


@app.get("/documents/{doc_id}")
def document(doc_id: str, s: Session = Depends(db)) -> dict[str, Any]:
    row = s.execute(text("SELECT * FROM documents WHERE document_id=:id"), {"id": doc_id}).mappings().first()
    if not row:
        raise HTTPException(404, "document not found")
    return dict(row)


@app.get("/documents/{doc_id}/artifacts")
def artifacts(doc_id: str, typ: str | None = None, s: Session = Depends(db)) -> dict[str, Any]:
    q = "SELECT artifact_type, content_json, file_path, created_at FROM artifacts WHERE document_id=:doc"
    params = {"doc": doc_id}
    if typ:
        q += " AND artifact_type=:typ"
        params["typ"] = typ
    rows = s.execute(text(q + " ORDER BY created_at"), params).mappings().all()
    return {"document_id": doc_id, "artifacts": [dict(r) for r in rows]}


def parse_file(doc: dict[str, Any]) -> dict[str, Any]:
    path = Path(doc["file_path"])
    ext = doc["file_type"]
    title = Path(doc["file_name"]).stem
    if ext == "txt":
        text_body = path.read_text(encoding="utf-8", errors="ignore")
        sections = [{"heading": title, "content": text_body, "order": 1}]
        tables = []
    elif ext == "md":
        body = path.read_text(encoding="utf-8", errors="ignore")
        parts = re.split(r"(?m)^#+\s+", body)
        headings = re.findall(r"(?m)^#+\s+(.+)$", body)
        sections = []
        if headings:
            for i, h in enumerate(headings):
                content = parts[i + 1].split("\n", 1)[1] if "\n" in parts[i + 1] else parts[i + 1]
                sections.append({"heading": h.strip(), "content": content.strip(), "order": i + 1})
        else:
            sections = [{"heading": title, "content": body, "order": 1}]
        tables = []
    elif ext == "docx":
        from docx import Document
        d = Document(str(path))
        text_body = "\n".join(p.text for p in d.paragraphs if p.text.strip())
        sections = [{"heading": title, "content": text_body, "order": 1}]
        tables = []
        for ti, table in enumerate(d.tables, 1):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            tables.append({"name": f"table_{ti}", "headers": rows[0] if rows else [], "rows": rows[1:] if len(rows) > 1 else []})
    elif ext == "pdf":
        import fitz
        pdf = fitz.open(str(path))
        sections = [{"heading": f"Page {i+1}", "content": page.get_text().strip(), "order": i + 1} for i, page in enumerate(pdf)]
        tables = []
    elif ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sections = [{"heading": title, "content": "Spreadsheet document", "order": 1}]
        tables = []
        for ws in wb.worksheets:
            rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
            if rows:
                tables.append({"name": ws.title, "headers": rows[0], "rows": rows[1:]})
    else:
        raise ValueError(f"unsupported parser: {ext}")
    return {"document_id": str(doc["document_id"]), "title": title, "sections": sections, "tables": tables, "metadata": {"file_type": ext, "parser": f"{ext}_parser"}}


def classify(parsed: dict[str, Any]) -> dict[str, Any]:
    text_all = "\n".join(s.get("content", "") for s in parsed.get("sections", []))
    sensitive = []
    level = "internal"
    if re.search(r"(password|token|secret|recipe|restricted|機密|confidential)", text_all, re.I):
        level = "confidential"
    if re.search(r"(機台|tool|equipment|SOP|CIM|IMC)", text_all, re.I):
        level = max([level, "internal"], key=["public", "internal", "confidential", "restricted"].index)
    for m in re.finditer(r"\b[A-Z]{2,}-?\d{2,}\b", text_all):
        sensitive.append({"text": m.group(0), "type": "possible_internal_id", "suggested_mask": "[INTERNAL_ID]"})
    return {"document_id": parsed["document_id"], "confidentiality": level, "masking_required": bool(sensitive), "sensitive_spans": sensitive, "reason": "Rule-assisted MVP classification. Replace with LLM structured output in production.", "confidence": 0.72}


def existing_entities(s: Session) -> list[dict[str, Any]]:
    rows = s.execute(text("""
        SELECT e.entity_id, e.canonical_name, e.entity_type, COALESCE(json_agg(a.alias) FILTER (WHERE a.alias IS NOT NULL), '[]') AS aliases
        FROM entities e LEFT JOIN entity_aliases a ON e.entity_id=a.entity_id
        GROUP BY e.entity_id, e.canonical_name, e.entity_type
    """)).mappings().all()
    return [dict(r) for r in rows]


def resolve_entities(s: Session, parsed: dict[str, Any]) -> dict[str, Any]:
    text_all = "\n".join(sec.get("content", "") for sec in parsed.get("sections", []))
    candidates = sorted(set(re.findall(r"[A-Za-z][A-Za-z0-9_\-]{2,}|[\u4e00-\u9fff]{2,8}", text_all)))[:40]
    ents = existing_entities(s)
    output = []
    for mention in candidates:
        if len(mention) < 2:
            continue
        exact = None
        for e in ents:
            names = [e["canonical_name"]] + list(e.get("aliases") or [])
            if mention.lower() in [str(n).lower() for n in names]:
                exact = e
                break
        if exact:
            output.append({"mention": mention, "decision": "existing_entity", "suggested_entity_id": str(exact["entity_id"]), "suggested_canonical_name": exact["canonical_name"], "entity_type": exact["entity_type"], "confidence": 0.98, "reason": "Exact canonical/alias match.", "requires_review": False})
        elif mention.isupper() or mention.endswith("Flow") or mention in ["機台", "製程", "流程"]:
            output.append({"mention": mention, "decision": "new_entity_candidate", "suggested_entity_id": None, "suggested_canonical_name": mention, "entity_type": "domain_term", "confidence": 0.7, "reason": "Important-looking term with no existing match.", "requires_review": True})
    return {"document_id": parsed["document_id"], "entities": output}


@app.post("/documents/{doc_id}/process")
def process(doc_id: str, s: Session = Depends(db)) -> dict[str, Any]:
    doc = document(doc_id, s)
    if doc["status"] == "DUPLICATE":
        raise HTTPException(400, "duplicate document cannot be processed")
    try:
        set_status(s, doc_id, "PARSING")
        parsed = parse_file(doc)
        save_artifact(s, doc_id, "parsed_json", parsed)
        set_status(s, doc_id, "PARSED")

        set_status(s, doc_id, "CLASSIFYING")
        cls = classify(parsed)
        save_artifact(s, doc_id, "classification_result", cls)
        set_status(s, doc_id, "CLASSIFIED")

        set_status(s, doc_id, "ENTITY_RESOLVING")
        resolved = resolve_entities(s, parsed)
        save_artifact(s, doc_id, "entity_resolution_result", resolved)
        for item in resolved["entities"]:
            rs = "auto_bound" if item["decision"] == "existing_entity" and item["confidence"] >= 0.95 else "pending"
            s.execute(text("""
                INSERT INTO entity_resolution_results(result_id, document_id, mention, decision, suggested_entity_id, suggested_canonical_name, entity_type, confidence, reason, requires_review, review_status)
                VALUES (:id,:doc,:m,:d,:eid,:name,:typ,:conf,:reason,:req,:rs)
            """), {"id": str(uuid.uuid4()), "doc": doc_id, "m": item["mention"], "d": item["decision"], "eid": item["suggested_entity_id"], "name": item["suggested_canonical_name"], "typ": item["entity_type"], "conf": item["confidence"], "reason": item["reason"], "req": item["requires_review"], "rs": rs})
        s.commit()
        pending = s.execute(text("SELECT COUNT(*) FROM entity_resolution_results WHERE document_id=:doc AND review_status='pending'"), {"doc": doc_id}).scalar_one()
        set_status(s, doc_id, "ENTITY_REVIEW_PENDING" if pending else "ENTITY_REVIEWED")
        return {"document_id": doc_id, "status": "ENTITY_REVIEW_PENDING" if pending else "ENTITY_REVIEWED"}
    except Exception as e:
        set_status(s, doc_id, "FAILED", str(e))
        raise


@app.get("/entities")
def list_entities(s: Session = Depends(db)) -> list[dict[str, Any]]:
    return existing_entities(s)


@app.get("/documents/{doc_id}/entity-results")
def entity_results(doc_id: str, s: Session = Depends(db)) -> list[dict[str, Any]]:
    rows = s.execute(text("SELECT * FROM entity_resolution_results WHERE document_id=:doc ORDER BY created_at"), {"doc": doc_id}).mappings().all()
    return [dict(r) for r in rows]


@app.post("/documents/{doc_id}/entity-review")
def entity_review(doc_id: str, req: EntityReviewRequest, s: Session = Depends(db)) -> dict[str, Any]:
    for action in req.actions:
        row = s.execute(text("SELECT * FROM entity_resolution_results WHERE result_id=:id AND document_id=:doc"), {"id": action.result_id, "doc": doc_id}).mappings().first()
        if not row:
            raise HTTPException(404, f"result not found: {action.result_id}")
        entity_id = action.entity_id
        if action.action == "create_new_entity":
            entity_id = str(uuid.uuid4())
            s.execute(text("INSERT INTO entities(entity_id, canonical_name, entity_type) VALUES (:id,:name,:typ)"), {"id": entity_id, "name": action.canonical_name or row["mention"], "typ": action.entity_type or row["entity_type"] or "domain_term"})
        elif action.action == "approve_alias":
            if not entity_id:
                raise HTTPException(400, "entity_id required")
            s.execute(text("INSERT INTO entity_aliases(alias_id, entity_id, alias, source_document_id, confidence, approved_by, approved_at) VALUES (:id,:eid,:alias,:doc,:conf,:by,NOW())"), {"id": str(uuid.uuid4()), "eid": entity_id, "alias": action.alias or row["mention"], "doc": doc_id, "conf": row["confidence"], "by": req.reviewed_by})
        elif action.action == "bind_existing_entity":
            if not entity_id:
                raise HTTPException(400, "entity_id required")
        elif action.action == "reject":
            s.execute(text("UPDATE entity_resolution_results SET review_status='rejected' WHERE result_id=:id"), {"id": action.result_id})
            continue
        else:
            raise HTTPException(400, f"unsupported action: {action.action}")
        s.execute(text("UPDATE entity_resolution_results SET review_status='approved', suggested_entity_id=:eid WHERE result_id=:id"), {"eid": entity_id, "id": action.result_id})
    s.commit()
    pending = s.execute(text("SELECT COUNT(*) FROM entity_resolution_results WHERE document_id=:doc AND review_status='pending'"), {"doc": doc_id}).scalar_one()
    set_status(s, doc_id, "ENTITY_REVIEWED" if pending == 0 else "ENTITY_REVIEW_PENDING")
    save_artifact(s, doc_id, "entity_review_result", {"reviewed_by": req.reviewed_by, "actions": [a.model_dump() for a in req.actions]})
    return {"document_id": doc_id, "status": "ENTITY_REVIEWED" if pending == 0 else "ENTITY_REVIEW_PENDING"}


def approved_mapping(s: Session, doc_id: str) -> list[dict[str, Any]]:
    rows = s.execute(text("""
      SELECT r.mention, r.suggested_entity_id AS entity_id, e.canonical_name, e.entity_type
      FROM entity_resolution_results r LEFT JOIN entities e ON r.suggested_entity_id=e.entity_id
      WHERE r.document_id=:doc AND r.review_status IN ('approved','auto_bound')
    """), {"doc": doc_id}).mappings().all()
    return [dict(r) for r in rows]


def build_okf_obj(s: Session, doc_id: str) -> dict[str, Any]:
    doc = document(doc_id, s)
    parsed = latest_artifact(s, doc_id, "parsed_json")
    cls = latest_artifact(s, doc_id, "classification_result")
    mapping = approved_mapping(s, doc_id)
    body = "\n".join(sec["content"] for sec in parsed.get("sections", []))
    evidence_id = "evidence_001"
    return {
        "okf_id": str(uuid.uuid4()),
        "schema_version": "v1",
        "title": parsed.get("title") or doc["file_name"],
        "source": {"document_id": doc_id, "source_id": doc["source_id"], "file_name": doc["file_name"]},
        "confidentiality": cls["confidentiality"],
        "summary": body[:240] or "Parsed document converted into OKF candidate.",
        "concepts": [{"concept_id": "concept_001", "name": parsed.get("title") or doc["file_name"], "definition": body[:300], "evidence_refs": [evidence_id]}],
        "entities": [{"entity_id": str(m.get("entity_id")), "name": m.get("canonical_name") or m["mention"], "type": m.get("entity_type") or "domain_term", "aliases": [m["mention"]]} for m in mapping if m.get("entity_id")],
        "procedures": [{"procedure_id": "procedure_001", "name": "Document procedure", "steps": [{"order": i + 1, "description": sec["heading"]} for i, sec in enumerate(parsed.get("sections", [])[:8])], "evidence_refs": [evidence_id]}],
        "relations": [{"relation_id": "relation_001", "subject": parsed.get("title") or doc["file_name"], "predicate": "describes", "object": "document knowledge", "evidence_refs": [evidence_id]}],
        "evidence": [{"evidence_id": evidence_id, "section": parsed.get("sections", [{}])[0].get("heading", "document"), "text": body[:1000]}],
        "review_status": "pending",
    }


def validate_okf_obj(okf: dict[str, Any], doc_id: str) -> list[str]:
    errors = []
    required = ["okf_id", "schema_version", "title", "source", "confidentiality", "summary", "concepts", "entities", "procedures", "relations", "evidence", "review_status"]
    for key in required:
        if key not in okf:
            errors.append(f"missing field: {key}")
    if okf.get("source", {}).get("document_id") != doc_id:
        errors.append("source.document_id mismatch")
    if okf.get("confidentiality") not in ["public", "internal", "confidential", "restricted"]:
        errors.append("invalid confidentiality")
    evidence_ids = {e.get("evidence_id") for e in okf.get("evidence", [])}
    for collection in ["concepts", "procedures", "relations"]:
        for item in okf.get(collection, []):
            for ref in item.get("evidence_refs", []):
                if ref not in evidence_ids:
                    errors.append(f"missing evidence ref {ref} in {collection}")
    return errors


@app.post("/documents/{doc_id}/build-okf")
def build_okf(doc_id: str, s: Session = Depends(db)) -> dict[str, Any]:
    pending = s.execute(text("SELECT COUNT(*) FROM entity_resolution_results WHERE document_id=:doc AND review_status='pending'"), {"doc": doc_id}).scalar_one()
    if pending:
        raise HTTPException(400, "entity review incomplete")
    set_status(s, doc_id, "BUILDING_OKF")
    okf = build_okf_obj(s, doc_id)
    save_artifact(s, doc_id, "okf_candidate", okf)
    errors = validate_okf_obj(okf, doc_id)
    save_artifact(s, doc_id, "schema_validation_result", {"valid": not errors, "errors": errors})
    set_status(s, doc_id, "OKF_REVIEW_PENDING" if not errors else "SCHEMA_INVALID")
    return {"document_id": doc_id, "status": "OKF_REVIEW_PENDING" if not errors else "SCHEMA_INVALID", "okf_json": okf, "errors": errors}


@app.post("/documents/{doc_id}/validate-okf")
def validate_okf(doc_id: str, req: ValidateOKFRequest) -> dict[str, Any]:
    errors = validate_okf_obj(req.okf_json, doc_id)
    return {"valid": not errors, "errors": errors}


@app.post("/documents/{doc_id}/approve-okf")
def approve_okf(doc_id: str, req: ApproveOKFRequest, s: Session = Depends(db)) -> dict[str, Any]:
    errors = validate_okf_obj(req.okf_json, doc_id)
    if errors:
        raise HTTPException(400, {"valid": False, "errors": errors})
    okf_id = req.okf_json["okf_id"]
    okf = dict(req.okf_json)
    okf["review_status"] = "approved"
    s.execute(text("INSERT INTO okf_documents(okf_id, document_id, schema_version, title, confidentiality, summary, okf_json, review_status, approved_by, approved_at) VALUES (:id,:doc,:sv,:title,:conf,:summary,CAST(:json AS JSONB),'approved',:by,NOW())"), {"id": okf_id, "doc": doc_id, "sv": okf["schema_version"], "title": okf["title"], "conf": okf["confidentiality"], "summary": okf.get("summary"), "json": to_json(okf), "by": req.approved_by})
    for typ, items in [("summary", [okf.get("summary", "")]), ("concept", [c.get("definition", "") for c in okf.get("concepts", [])]), ("procedure", ["; ".join(step.get("description", "") for step in p.get("steps", [])) for p in okf.get("procedures", [])]), ("evidence", [e.get("text", "") for e in okf.get("evidence", [])])]:
        for content in items:
            if content:
                s.execute(text("INSERT INTO okf_chunks(chunk_id, okf_id, chunk_type, content, metadata_json) VALUES (:id,:okf,:typ,:content,CAST(:meta AS JSONB))"), {"id": str(uuid.uuid4()), "okf": okf_id, "typ": typ, "content": content, "meta": to_json({"title": okf["title"], "confidentiality": okf["confidentiality"], "source_id": okf["source"]["source_id"], "document_id": doc_id})})
    s.commit()
    save_artifact(s, doc_id, "approved_okf", okf)
    set_status(s, doc_id, "READY")
    return {"document_id": doc_id, "okf_id": okf_id, "status": "READY"}


@app.post("/documents/{doc_id}/reject-okf")
def reject_okf(doc_id: str, req: RejectOKFRequest, s: Session = Depends(db)) -> dict[str, Any]:
    set_status(s, doc_id, "OKF_REVIEW_REJECTED", req.reason)
    return {"document_id": doc_id, "status": "OKF_REVIEW_REJECTED"}


@app.get("/search")
def search(q: str, s: Session = Depends(db)) -> dict[str, Any]:
    rows = s.execute(text("""
      SELECT od.okf_id, c.chunk_id, od.title, c.chunk_type, c.content, c.metadata_json,
             ts_rank_cd(to_tsvector('simple', c.content), plainto_tsquery('simple', :q)) AS score
      FROM okf_chunks c JOIN okf_documents od ON c.okf_id=od.okf_id
      WHERE od.review_status='approved'
        AND (to_tsvector('simple', c.content) @@ plainto_tsquery('simple', :q) OR c.content ILIKE :like)
      ORDER BY score DESC NULLS LAST
      LIMIT 20
    """), {"q": q, "like": f"%{q}%"}).mappings().all()
    return {"query": q, "results": [dict(r) for r in rows]}
